from pathlib import Path
import shutil
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

BASE = Path(r"D:\path\to\morchella_microbiome\analysis\manuscript_v1")
SRC = BASE / "Applied_Soil_Ecology_顶刊风格图文整合完整初稿_v6_2026-08-22.docx"
OUT = BASE / "Applied_Soil_Ecology_顶刊风格图文整合完整初稿_v9_Figure7终版_2026-08-22.docx"
FIG = BASE / "figures_v3" / "Figure_7_FUNGuild_trophic_modes_v8.png"

shutil.copy2(SRC, OUT)
doc = Document(OUT)

def find_exact(text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"Paragraph not found: {text}")

def find_prefix(prefix):
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise ValueError(f"Paragraph prefix not found: {prefix}")

def add_before(anchor, text, style="Normal"):
    p = anchor.insert_paragraph_before(text, style=style)
    return p

def add_after(anchor, text, style="Normal"):
    p = doc.add_paragraph(text, style=style)
    anchor._p.addnext(p._p)
    return p

# Abstract: add only the result that changes the paper-level interpretation.
abstract = doc.paragraphs[17]
needle = "Community restructuring included increases in Patescibacteria, Bacteroidota and Mortierellomycota and decreases in Actinomycetota and Ascomycota."
addition = (" Community-level separation remained positive after each greenhouse was omitted. "
            "FUNGuild prediction further indicated a directionally consistent increase in the saprotrophic fraction of annotated fungal reads after cultivation "
            "(8 of 9 pairs; exact P = 0.0214; false-discovery-rate-adjusted q = 0.0643).")
if needle not in abstract.text:
    raise ValueError("Abstract insertion anchor missing")
abstract.text = abstract.text.replace(needle, needle + addition)

# Methods: insert a dedicated reproducible module and renumber statistical reporting.
stat_head = find_exact("2.11. Statistical reporting and reproducibility")
stat_head.text = "2.12. Statistical reporting and reproducibility"
new_head = add_before(stat_head, "2.11. Exact paired sensitivity analysis and fungal guild prediction", "Heading 2")
add_before(stat_head,
    "As a robustness analysis of cultivation-stage separation, a paired distance-ratio statistic (reported as a pseudo-F value) was recalculated for bacterial and fungal Bray–Curtis matrices. Statistical significance was evaluated by enumerating all 2^9 within-pair stage-label swaps (512 permutations). The statistic was then recalculated after omitting each greenhouse in turn; these leave-one-greenhouse-out values were treated as sensitivity estimates rather than independent hypothesis tests. Candidate-genus centered-log-ratio directions were additionally recalculated using count-scale pseudocounts of 0.1, 0.5 and 1.",
    "Normal")
add_before(stat_head,
    "Fungal ASVs were annotated against the FUNGuild database (Nguyen et al., 2016). The strict analysis retained only Probable and Highly Probable assignments, whereas the inclusive analysis additionally retained Possible assignments. When an ASV had multiple trophic modes or guilds, its count was allocated fractionally across the assigned categories. Relative abundance was calculated using either all ITS reads or only FUNGuild-annotated reads as the denominator. For the nine spatial pairs, after-minus-before differences were tested using all 2^9 paired sign flips. Benjamini–Hochberg correction was applied separately within each confidence-set, ontology and denominator family. Leave-one-greenhouse-out mean effects were used to evaluate directional robustness. FUNGuild outputs were interpreted as database-predicted ecological categories rather than measured metabolic functions or activities.",
    "Normal")

# Results: report the robustness analysis and functional prediction separately from interpretation.
table_head = find_exact("Table 1. Frozen candidate genera and external evidence classification")
add_before(table_head, "3.8. Exact paired sensitivity analyses supported restructuring and indicated a saprotrophic shift", "Heading 2")
add_before(table_head,
    "Exact enumeration of all 512 within-pair label swaps supported cultivation-stage separation for both bacterial and fungal communities (16S pseudo-F = 1.1135, P = 0.00585; ITS pseudo-F = 1.1136, P = 0.00585). The effect remained positive after each greenhouse was omitted: pseudo-F values ranged from 1.1055 to 1.1698 for 16S and from 1.0636 to 1.3047 for ITS. Candidate-genus directions were also invariant across pseudocounts of 0.1, 0.5 and 1 for Terrimonas (9/9 positive pairs), Chitinophaga (9/9), Pseudarthrobacter (9/9 negative), Morchella (9/9 negative), Mortierella (9/9 positive) and Alternaria (9/9 negative). Gemmata retained a negative direction in eight of nine pairs and Botryotrichum a positive direction in seven of nine pairs at every pseudocount.",
    "Normal")
add_before(table_head,
    "FUNGuild assigned an ecological guild to 1,641 of 8,056 ITS ASVs (20.37%). Under the strict confidence set and the annotated-read denominator, the saprotrophic fraction increased by 13.0 percentage points after cultivation, with the same direction in eight of nine pairs (exact P = 0.0214; BH q = 0.0643; Fig. 7a). The inclusive confidence set produced a similar increase of 10.8 percentage points (8/9 pairs; P = 0.0214; q = 0.0643). Positive saprotroph effects persisted after omitting each greenhouse under both confidence definitions (Fig. 7b). Pathotroph and symbiotroph fractions did not show FDR-supported changes. Several individual guilds had nominal P values below 0.05, but none reached q < 0.10; these signals were therefore retained as exploratory only.",
    "Normal")

# Discussion: separate direct prediction from ecological interpretation and mechanism speculation.
lim_head = find_exact("4.7. Limitations and applied implications")
lim_head.text = "4.8. Limitations and applied implications"
add_before(lim_head, "4.7. Predicted trophic-mode redistribution suggests greater saprotrophic representation", "Heading 2")
add_before(lim_head,
    "The abundance-weighted FUNGuild analysis added a functional-category layer to the taxonomic results. The direct result is limited to a higher predicted saprotrophic share among annotated fungal reads after cultivation, supported by eight of nine paired positions and by leave-one-greenhouse-out sensitivity estimates. This pattern is compatible with the observed increase in Mortierella and the supporting increase in Botryotrichum, both of which received saprotrophic components in the database annotation. It also provides a coherent ecological interpretation for a post-cultivation environment containing altered organic substrates and fungal or plant residues.",
    "Normal")
add_before(lim_head,
    "This interpretation should not be converted into a measured decomposition mechanism. Only 20.37% of ITS ASVs received a FUNGuild assignment, annotations were transferred from taxonomic identity, and relative abundance among annotated reads is sensitive to both the annotated subset and compositional closure. The result therefore indicates redistribution of predicted trophic categories, not higher extracellular-enzyme activity, decomposition rate or carbon mineralization. Metagenomic or metatranscriptomic profiling, enzyme assays and substrate-use experiments would be required to test whether the predicted saprotrophic shift corresponds to realized function.",
    "Normal")

# Conclusion: integrate the new evidence without elevating prediction to direct function.
conclusion = find_prefix("Paired full-length 16S and ITS profiling showed")
needle2 = "Community-level reorganization was therefore more transferable than any universal genus-level direction."
extra2 = (" Exact paired sensitivity analyses supported the stability of the community and candidate-genus directions, while FUNGuild suggested an increased saprotrophic share of the annotated fungal community after cultivation.")
if needle2 not in conclusion.text:
    raise ValueError("Conclusion insertion anchor missing")
conclusion.text = conclusion.text.replace(needle2, needle2 + extra2)

# Add the canonical FUNGuild reference in alphabetical position.
near_ref = find_prefix("Nearing, J.T.")
add_before(near_ref,
    "Nguyen, N.H., Song, Z., Bates, S.T., Branco, S., Tedersoo, L., Menke, J., Schilling, J.S., Kennedy, P.G., 2016. FUNGuild: An open annotation tool for parsing fungal community datasets by ecological guild. Fungal Ecol. 20, 241–248. https://doi.org/10.1016/j.funeco.2015.06.006.",
    "Normal")

# Insert Figure 7 and its self-contained legend immediately before Fig. S1.
fig6_caption = find_prefix("Fig. 6.")
img_p = doc.add_paragraph()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_p.paragraph_format.page_break_before = True
img_p.paragraph_format.keep_with_next = True
img_p.add_run().add_picture(str(FIG), width=Inches(6.65))
fig6_caption._p.addnext(img_p._p)

legend_text = (
    "Fig. 7. Predicted fungal trophic-mode redistribution after morel cultivation. "
    "(a) Mean composition of pathotroph, saprotroph and symbiotroph categories before and after cultivation under the strict confidence set and the annotated-read denominator. Ribbon widths and adjacent percentages represent stage means. "
    "(b) Violin distributions and embedded box plots summarize the after-minus-before changes across the nine spatial pairs for each trophic mode. Boxes show the interquartile range and median, whiskers extend to values within 1.5 times the interquartile range, and labels report the number of pairs with positive changes. "
    "(c) Sensitivity matrix of mean paired effects across strict and inclusive confidence sets and across annotated-read and all-read denominators. Cell values are mean after-minus-before changes in percentage points; pink and blue indicate negative and positive effects, respectively. The q values shown for the annotated saprotroph cells were obtained using Benjamini–Hochberg correction within each confidence-set, ontology and denominator family. n = 9 spatial pairs nested within three greenhouses. Counts from ASVs assigned to multiple categories were allocated fractionally. Exact P values were obtained from all 2^9 within-pair sign flips; the strict annotated-read saprotroph result was positive in eight of nine pairs (P = 0.0214, q = 0.0643). FUNGuild categories are database predictions and do not represent measured metabolic activity."
)
cap_p = doc.add_paragraph(legend_text, style="Figure legend")
cap_p.paragraph_format.keep_together = True
img_p._p.addnext(cap_p._p)

# Ensure the supplementary figure starts on a fresh page after the new legend.
s1_image = cap_p._p.getnext()
if s1_image is not None:
    from docx.text.paragraph import Paragraph
    try:
        Paragraph(s1_image, cap_p._parent).paragraph_format.page_break_before = True
    except Exception:
        pass

doc.save(OUT)
print(OUT)

