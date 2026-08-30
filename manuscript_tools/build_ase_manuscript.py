from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

ROOT = Path(r"D:\path\to\morchella_microbiome\analysis\manuscript_v1")
OUT = ROOT / "Applied_Soil_Ecology_顶刊风格图文整合完整初稿_v6_2026-08-22.docx"
HIGHLIGHTS = ROOT / "Applied_Soil_Ecology_Highlights_v2_2026-08-22.docx"
SUMMARY = ROOT / "羊肚菌微生物组_导师汇报双总结_v1_2026-08-22.docx"

def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)

def line_numbering(section):
    sectPr = section._sectPr
    ln = sectPr.find(qn("w:lnNumType"))
    if ln is None:
        ln = OxmlElement("w:lnNumType")
        sectPr.append(ln)
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")

def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1.15)
    sec.right_margin = Inches(1.15)
    line_numbering(sec)
    add_page_number(sec.footer.paragraphs[0])
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True
    for sname, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)]:
        st = styles[sname]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.line_spacing = 2
        st.paragraph_format.space_before = Pt(6)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.keep_with_next = True
    if "Figure legend" not in styles:
        st = styles.add_style("Figure legend", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
        st.font.size = Pt(10)
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.space_after = Pt(6)

def p(doc, text="", bold_lead=None, italic_terms=()):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = 2
    if bold_lead and text.startswith(bold_lead):
        r = para.add_run(bold_lead); set_font(r, bold=True)
        text = text[len(bold_lead):]
    # limited scientific-name formatting
    cursor = 0
    terms = sorted(italic_terms, key=lambda x: text.find(x) if x in text else 10**9)
    for term in terms:
        idx = text.find(term, cursor)
        if idx < 0: continue
        if idx > cursor:
            r = para.add_run(text[cursor:idx]); set_font(r)
        r = para.add_run(term); set_font(r, italic=True)
        cursor = idx + len(term)
    if cursor < len(text):
        r = para.add_run(text[cursor:]); set_font(r)
    return para

def heading(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    r = para.add_run(text); set_font(r, "Arial", 14 if level == 1 else 12, bold=True)
    return para

def add_title_page(doc):
    q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before = Pt(72); q.paragraph_format.space_after = Pt(24)
    r = q.add_run("Cultivation-stage-associated bacterial–fungal community restructuring and context-dependent genus responses in Morchella production soils")
    set_font(r, "Arial", 17, bold=True)
    for line in [
        "",
        "[Author 1]a, [Author 2]b, [Author 3]a, [Corresponding author]a,*",
        "",
        "a [Department, Institution, City, Postal code, Country]",
        "b [Department, Institution, City, Postal code, Country]",
        "",
        "* Corresponding author: [Name]",
        "E-mail: [email address]",
        "Postal address: [full postal address]",
        "Telephone: [telephone number]",
        "",
        "Article type: Research Paper",
        "Target journal: Applied Soil Ecology",
        "Manuscript status: complete working draft; author metadata and wet-lab details pending confirmation"
    ]:
        para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = para.add_run(line); set_font(rr, size=11, italic=line.startswith("Manuscript status"))
    doc.add_page_break()

def add_abstract(doc):
    heading(doc, "Abstract", 1)
    abstract = ("Stable morel production depends on soil conditions, yet bacterial and fungal responses during cultivation and the transferability of candidate taxa across production systems remain uncertain. We profiled 24 production-soil samples using PacBio full-length 16S rRNA gene and internal transcribed spacer amplicons. The core comparison comprised nine spatially matched before–after pairs nested within three greenhouses. Cultivation stage explained 15.2% of bacterial Bray–Curtis variation and approximately 14–15% of fungal variation. Bacterial and fungal distance matrices were strongly correlated (Mantel r = 0.590 for the 18 paired samples), whereas paired changes in Shannon diversity and observed amplicon sequence variants were not correlated. Community restructuring included increases in Patescibacteria, Bacteroidota and Mortierellomycota and decreases in Actinomycetota and Ascomycota. Eight genera were frozen before external testing after evaluation of compositional effect size, paired-direction consistency, greenhouse-level direction and sequence-level taxonomic support. Reanalysis of three public BioProjects provided partial support for Morchella and Pseudarthrobacter, context-dependent evidence for four genera, opposing evidence for Gemmata and insufficient evidence for Alternaria. Thus, community-level restructuring was more transferable than the direction of individual genera. Because cultivation stage was confounded with sequencing batch and samples were nested within only three greenhouses, the findings identify design-bounded candidates rather than causal drivers or universal biomarkers.")
    p(doc, abstract)
    p(doc, "Keywords: soil microbiome; edible fungi; full-length amplicon; paired sampling; community ecology; external validation; compositional data")

def add_intro(doc):
    heading(doc, "1. Introduction", 1)
    paras = [
    "Morels (Morchella spp.) are high-value edible ascomycetes whose commercial production has expanded rapidly, particularly in field and greenhouse systems. Production nevertheless remains unstable, with recurrent yield loss, soil deterioration under repeated cultivation and outbreaks of soil- or substrate-associated disease. The microbial community surrounding morel mycelia is central to this production environment because bacteria and fungi mediate decomposition, nutrient turnover and the transformation of exogenous nutrient bags and crop residues. Temporal studies have further shown that bacterial and fungal assemblages change across mycelial growth, primordium formation and fruiting, linking microbial succession to the changing resource environment of cultivation (Longley et al., 2019; Tan et al., 2021; Zhang et al., 2023).",
    "Interpreting these shifts is difficult in production soils. Microbial composition varies with location, soil physicochemistry, cultivation history, management and sampling time. Consequently, an unpaired comparison between cultivated and uncultivated plots can conflate the cultivation response with pre-existing spatial heterogeneity. Continuous-cropping studies have reported lower productivity and altered soil microbial communities in Morchella systems, but when cultivation history corresponds to different greenhouses or locations, the apparent year effect cannot be separated from greenhouse-specific conditions (Liu et al., 2022). Spatially matched before–after sampling can reduce baseline heterogeneity, although it cannot by itself eliminate temporal or technical confounding.",
    "A second challenge concerns the level at which transferability is evaluated. Most amplicon studies describe alpha diversity, beta diversity and differentially abundant taxa within one experiment. These summaries are useful, but bacterial and fungal datasets are often considered separately, even when obtained from the same soil. Moreover, taxa selected from a discovery dataset may not show the same direction in another production system because amplification region, sequencing platform, growth stage and management regime differ. Microbiome differential-abundance methods can themselves return substantially different candidate sets, and relative abundance does not directly measure absolute population size (Knight et al., 2018; Morton et al., 2019; Lin and Peddada, 2020; Nearing et al., 2022). Candidate selection should therefore incorporate effect magnitude, within-pair direction, independent sampling-unit structure and taxonomic reliability before external data are examined.",
    "Here, we combined full-length 16S rRNA gene and ITS amplicon sequencing for soils collected from identical spatial positions before and after morel cultivation. We asked whether bacterial and fungal communities underwent coordinated structural reorganization, whether their within-community diversity changed synchronously and whether directionally stable candidate genera from the local dataset transferred to independent public datasets. Candidate genera were frozen before public-data reanalysis, and external evidence was classified as supportive, context dependent, opposing or insufficient. A small white-mold point comparison and soil physicochemical measurements were retained as exploratory context rather than as independent disease or environmental-driver tests. This design provides a transparent assessment of which conclusions are robust at the community level and where genus-level interpretation becomes system dependent."
    ]
    for x in paras: p(doc, x, italic_terms=("Morchella spp.",))

def add_methods(doc):
    heading(doc, "2. Materials and methods", 1)
    sections = [
    ("2.1. Study design and soil sampling", [
    "The primary cultivation-stage comparison included three commercial morel greenhouses. Three spatial locations were selected within each greenhouse. Soil was collected from each location before cultivation (M) and after cultivation (A), yielding nine spatially matched pairs and 18 samples. The three greenhouses had histories of two, three and four years of continuous morel cultivation, respectively. Because each cultivation history corresponded to only one greenhouse, cultivation duration was completely confounded with greenhouse identity and was not treated as an independently estimable exposure. Soil type, morel cultivar, management regime and sampling period were reported as consistent across the three greenhouses.",
    "For exploratory disease context, one visually healthy location and one white-mold-affected location were sampled simultaneously within a single greenhouse. Three spatial subsamples were obtained from each location, yielding six additional samples. These subsamples characterize two local points and were not treated as six independent disease-status replicates. In total, 24 biological soil samples were profiled using both bacterial 16S and fungal ITS markers. [AUTHOR INPUT NEEDED: sampling date, geographic coordinates or administrative location, soil depth, distance between locations, mass collected, homogenization procedure, storage temperature and time to DNA extraction.]"
    ]),
    ("2.2. Soil physicochemical measurements", [
    "Eight soil properties were quantified for the nine before–after sampling positions: pH, total nitrogen, total phosphorus, total potassium, available nitrogen, available phosphorus, available potassium and organic matter [confirm exact eighth variable from the analytical report]. Concentrations reported by the testing laboratory in g kg−1 were retained in their original units, with unit conversion applied only where explicitly indicated in the source workbook. For paired analyses, the change in each property was calculated as A minus M at the same spatial position. [AUTHOR INPUT NEEDED: analytical laboratory, extraction and detection methods for each property, instrument models, method standards and quality-control procedures.]"
    ]),
    ("2.3. DNA extraction, amplification and PacBio sequencing", [
    "According to the sequencing-provider materials currently available, total soil DNA was extracted using a cetyltrimethylammonium bromide protocol, checked by 1% agarose gel electrophoresis and quantified fluorometrically using Qubit. This description must be confirmed against the project-specific laboratory record because the accompanying writing package also contains generic protocols that were not necessarily used for these samples.",
    "The full-length bacterial 16S rRNA gene was amplified using primers AGRGTTYGATYMTGGCTCAG and AAGTCGTAACAAGGTARCY. The fungal ITS region was amplified using TACACACCGCCCGTCG and CCTSCSCTTANTDATATGC. Purified amplicons were quantified, pooled in equimolar proportions and sequenced as PacBio circular consensus sequences. [AUTHOR INPUT NEEDED: PCR reaction volume, template input, primer concentrations, polymerase and reagent manufacturers, thermocycling conditions, library-preparation kit and version, PacBio instrument model, sequencing chemistry, and CCS-generation software and thresholds.] Generic Illumina PE250 parameters from the provider’s writing package were not substituted for these missing PacBio-specific records."
    ]),
    ("2.4. Read integrity, primer removal and ASV inference", [
    "Raw FASTQ integrity was verified against provider-supplied MD5 checksums. Primers were identified and removed with Cutadapt v5.2 in linked-adapter mode, requiring both terminal primers, allowing reverse-complement orientation, using a maximum error rate of 0.10 and discarding reads without both primer matches. For ITS, the reverse-complement sequence GCATATHANTAAGSGSAGG was used to recognize the 3′ terminal primer orientation corresponding to the supplied primer CCTSCSCTTANTDATATGC.",
    "Bacterial and fungal reads were processed separately in R v4.4.2 with DADA2 v1.34.0 (Callahan et al., 2016). Filtering parameters were maxN = 0, maxEE = 2, truncQ = 2 and rm.phix = FALSE. Retained length ranges were 1,200–1,800 bp for 16S and 400–1,500 bp for ITS. Samples were divided according to the original M, A and disease/healthy sequencing batches. For error learning, up to 10,000 filtered reads per sample were balanced within each batch. PacBioErrfun was used for the error model, ASVs were inferred with pseudo-pooling, and BAND_SIZE = 32 was additionally specified for 16S. Batch-specific sequence tables were merged and chimeras were removed using the consensus method. Processing was performed separately by marker; 16S and ITS read counts were never merged."
    ]),
    ("2.5. Taxonomic assignment and sequence-level reliability", [
    "Bacterial ASVs were classified with dada2::assignTaxonomy against the SILVA v138.2 DADA2 training set using minBoot = 80 and tryRC = TRUE. Species names were added only for exact matches using addSpecies. Chloroplast and mitochondrial ASVs were removed before downstream analysis.",
    "ITS ASVs were queried using BLASTN against the UNITE general FASTA release v10.0 archived on 19 February 2025. Best-hit identity, query coverage, E-value, bitscore, taxonomic rank and Species Hypothesis identifier were retained. Genus-level reporting generally required ≥97% identity, ≥80% query coverage and agreement among highest-scoring genus assignments. Analyses were performed both on all ITS ASVs and on the subset assigned to Fungi to evaluate sensitivity to unclassified and non-fungal reads. Candidate-genus reliability was reviewed using 16S bootstrap support or ITS alignment identity, coverage and ASV composition; species-level interpretations were not made where support was insufficient."
    ]),
    ("2.6. Alpha and beta diversity", [
    "Observed ASV richness and Shannon diversity were calculated separately for each marker. The primary bacterial diversity analysis used random subsampling without replacement to 1,500 reads per sample, which retained all 24 samples; sensitivity analyses used depths of 1,800 and 2,500 reads. ITS data were subsampled to 16,000 reads per sample. Cultivation-stage comparisons used the spatial position as the paired unit.",
    "Bray–Curtis dissimilarity was the primary beta-diversity metric, with binary Jaccard distance used as a bacterial sensitivity analysis. Principal coordinates analysis (PCoA) was used for visualization. The M/A association was tested with permutational multivariate analysis of variance (PERMANOVA) using permutations restricted within paired positions. Permutational analysis of multivariate dispersions (PERMDISP) assessed whether apparent group separation could be attributed to unequal within-group dispersion. Because the paired samples were nested in three greenhouses and cultivation stage coincided with sequencing batch, permutation-derived P values were interpreted as exploratory rather than as unconfounded causal tests."
    ]),
    ("2.7. Community composition and candidate-genus freezing", [
    "Phylum-level composition was summarized as sample relative abundance. For a descriptive overview of genus-level shifts, taxa were required to occur in at least nine M/A samples and were ranked by the product of the absolute log2 ratio of group means and the larger group mean relative abundance. This ranking was used only to display compositionally prominent changes and was not treated as a false-discovery-rate-controlled differential-abundance list.",
    "Candidate genera were selected using a pre-specified evidence combination: centered-log-ratio effect, direction across the nine spatial pairs, direction of greenhouse-level means, abundance, ASV composition and taxonomic reliability. Seven primary candidates and one supporting candidate were frozen before any public-dataset result was inspected. Public evidence therefore could change the evidence grade but could not retrospectively add or remove candidates. Because the measurements were compositional, candidate shifts were described as relative-abundance changes rather than absolute increases or decreases."
    ]),
    ("2.8. Cross-domain analysis", [
    "Bray–Curtis matrices were computed independently for 16S and ITS after marker-specific processing. A Mantel test assessed concordance between the bacterial and fungal sample-distance matrices. The test was conducted for all 24 samples and separately for the 18 M/A samples. To distinguish community-level coordination from within-community diversity responses, paired A-minus-M changes in Shannon diversity and observed ASVs were calculated for each marker and compared using Spearman correlation across the nine spatial positions."
    ]),
    ("2.9. External evaluation using public amplicon datasets", [
    "Three morel-associated BioProjects were included: PRJNA935967 (ITS), PRJNA993383 (16S) and PRJNA822526 (16S and ITS). Each project was processed independently according to its primers and sequencing platform, including project-specific quality control, DADA2 inference and taxonomic assignment. ASVs and count tables were not merged across projects or with the local dataset. External evaluation was restricted to the eight frozen genera and compared within-project effect directions, detectability and consistency with the pre-specified local direction.",
    "PRJNA935967 and PRJNA993383 originated from the same experimental setting and were therefore counted jointly as one external evidence system. PRJNA822526 represented a second system; however, its three files per group were operational splits of a pooled composite sample. Its biological sample size was therefore treated as one per group, and inferential P values were not calculated. External findings were classified as partially supported, context dependent, externally opposed or insufficient, rather than being described as an independent validation cohort."
    ]),
    ("2.10. Exploratory soil and white-mold analyses", [
    "For the nine M/A pairs, Spearman correlations were calculated between paired changes in candidate-genus centered-log-ratio abundance and paired changes in each soil property. Benjamini–Hochberg adjustment controlled the false discovery rate across genus–soil tests. The white-mold comparison summarized alpha diversity for one healthy and one affected point. No disease-status inferential test was performed because the three samples per point were spatial subsamples and did not provide independent healthy or diseased locations."
    ]),
    ("2.11. Statistical reporting and reproducibility", [
    "Effect sizes, paired-direction counts and sensitivity analyses were emphasized alongside P values. Unless otherwise stated, tests were two-sided and exact sample sizes are given in the text or figure legends. Computational software versions, scripts, intermediate tables, session information and frozen candidate evidence were retained in the project analysis archive. [AUTHOR INPUT NEEDED: permanent code repository and sequence-data accession. Raw sequence data should be deposited in a public repository before submission, or the accession should be stated as pending during peer review if permitted.]"
    ])]
    for h, paras in sections:
        heading(doc, h, 2)
        for x in paras: p(doc, x, italic_terms=("Morchella", "Morchella spp."))

def add_results(doc):
    heading(doc, "3. Results", 1)
    sections = [
    ("3.1. Paired sampling and full-length amplicon data supported cross-domain comparison", [
    "The dataset comprised 24 production-soil samples, each profiled with full-length 16S and ITS amplicons. The core M/A comparison contained nine matched positions distributed among three greenhouses, whereas the six white-mold-context samples represented three subsamples from each of two points (Fig. 1a,b). All 48 raw FASTQ files passed MD5 verification. Primer recognition and removal retained 99.75% of 16S reads and 99.44% of ITS reads.",
    "After denoising and chimera removal, the ITS dataset contained 8,056 ASVs and 942,492 reads, with all 24 samples retained. UNITE assigned 2,510 ASVs to Fungi, representing 76.77% of ITS reads. Bray–Curtis matrices from all ITS ASVs and the fungal-only subset were nearly identical (Spearman ρ = 0.9852), indicating that the principal between-sample structure was not driven solely by uncertain high-rank assignments. Before bacterial taxonomic filtering, the 16S dataset contained 9,752 ASVs and 193,540 reads. Removal of chloroplast and mitochondrial sequences retained 9,477 ASVs and 188,611 reads (97.45% of sequences; Fig. 1c,d)."
    ]),
    ("3.2. Cultivation stage was associated with bacterial and fungal community restructuring", [
    "Bacterial communities separated between M and A samples in Bray–Curtis ordination while preserving the matched-position trajectories (Fig. 2a). At the primary rarefaction depth of 1,500 reads, restricted PERMANOVA attributed 15.2% of bacterial community variation to M/A stage (R² = 0.152, exploratory P = 0.0039). Binary Jaccard analysis gave the same direction (R² = 0.117, exploratory P = 0.0039). Dispersion did not differ detectably between M and A for either Bray–Curtis (PERMDISP P = 0.859) or Jaccard distance (P = 0.532), arguing against dispersion alone as the source of group separation.",
    "The bacterial result was stable to sequencing-depth choice. At 1,800 and 2,500 reads, Bray–Curtis effect sizes were R² = 0.162 and 0.176, with exploratory P = 0.0078 and 0.0313, respectively. Alpha-diversity changes were more modest: observed ASVs increased by a mean of 53.8 and Shannon diversity by 0.206, with increases in six of nine pairs for both indices (Fig. 2c).",
    "Fungal composition also changed with cultivation stage (Fig. 2b). Analyses of all ITS ASVs and the fungal-only subset produced similar M/A effect sizes of approximately R² = 0.14–0.15. Within the fungal subset, Shannon diversity increased in six of nine pairs, with a mean change of +0.942. Observed fungal ASVs also increased in six pairs, but the mean change was −13.3 because one greenhouse showed a comparatively large negative shift (Fig. 2d). These associations combine cultivation-stage, batch and greenhouse-linked variation and therefore do not identify a cultivation effect independent of those factors."
    ]),
    ("3.3. Dominant bacterial and fungal groups underwent compositional replacement", [
    "The M/A difference involved broad replacement among abundant phyla (Fig. 3a,b). Mean bacterial relative abundance increased from 12.22% to 28.64% for Patescibacteria and from 11.70% to 18.41% for Bacteroidota. In contrast, Actinomycetota decreased from 18.28% to 1.85% and Planctomycetota from 13.32% to 9.63%. Pseudomonadota remained abundant in both stages (24.66% in M and 23.49% in A).",
    "Among fungi, Ascomycota decreased from 65.90% to 49.61%, whereas Mortierellomycota increased from 2.87% to 10.91%. Unclassified or higher-rank-only reads increased from 24.40% to 33.68%. Abundance-weighted genus ranking identified several high-magnitude increases and decreases in each marker (Fig. 3c,d). This ranking provided a community overview; it did not imply that the displayed taxa passed multiplicity-adjusted differential-abundance tests."
    ]),
    ("3.4. Bacterial and fungal beta structures were coordinated, whereas alpha-diversity responses were not", [
    "Across all 24 samples, bacterial and fungal Bray–Curtis matrices were positively correlated (Mantel r = 0.587, P = 0.0001). The relationship remained similar when restricted to the 18 paired M/A samples (r = 0.590, P = 0.0001; Fig. 4a). Thus, samples that were compositionally dissimilar in the bacterial dataset tended also to be dissimilar in the fungal dataset.",
    "This beta-structure coordination did not extend to paired alpha-diversity changes. Across the nine spatial pairs, bacterial and fungal changes were uncorrelated for Shannon diversity (Spearman ρ = −0.167, P = 0.668) and observed ASVs (ρ = −0.159, P = 0.683; Fig. 4b–d). A bacterial richness or evenness increase at a location therefore did not predict the corresponding fungal alpha-diversity direction."
    ]),
    ("3.5. Eight frozen candidate genera showed stable or near-stable local compositional shifts", [
    "Before public datasets were examined, seven primary candidate genera and one supporting genus were frozen using compositional effect, pairwise direction, greenhouse-level direction, ASV composition and classification support (Table 1; Fig. 5). Read-weighted genus-level bootstrap support for the four bacterial candidates ranged from 96.51% to 100%. Weighted UNITE identity for fungal candidates ranged from 97.51% to 99.84%, supporting genus-level but not general species-level interpretation.",
    "Among fungi, Morchella decreased from a mean relative abundance of 29.90% before cultivation to no detection after cultivation, with decreases in all nine pairs and all three greenhouse means. One dominant ASV accounted for 97.6% of reads assigned to this genus. Mortierella increased from 2.41% to 15.46% in all nine pairs and was represented by multiple ASVs. Alternaria decreased from 0.353% to 0.020% in all nine pairs. Botryotrichum increased overall from 2.25% to 15.26%, but only seven pairs increased and the four-year-history greenhouse showed the opposite mean direction; it was therefore retained as a supporting rather than primary candidate.",
    "The four bacterial candidates also showed strong within-pair consistency. Terrimonas increased from 0.467% to 1.574%, and Chitinophaga increased from 0.051% to 0.989%; both increased in all nine pairs. Gemmata decreased from 0.564% to 0.132% in eight pairs, while Pseudarthrobacter decreased from 8.955% to 0.159% in all nine pairs. All four genera had concordant mean directions across the three greenhouses. These are relative compositional shifts and do not establish equivalent changes in absolute abundance or metabolic activity."
    ]),
    ("3.6. Public datasets indicated partial transferability and substantial context dependence", [
    "Three BioProjects were processed independently and evaluated only after the candidate list had been frozen. PRJNA935967 and PRJNA993383 formed the first external system, while PRJNA822526 formed the second. Because the external contrasts did not exactly reproduce the local before–after design, they were treated as directional stress tests rather than pooled replication.",
    "Morchella and Pseudarthrobacter received partial external support (Fig. 6). In the first system, two of three Morchella stage contrasts supported the pre-specified decrease and one was indeterminate. Both contrasts in the second system trended downward but had weak design support. Pseudarthrobacter decreased in all three first-system contrasts and was not detected in the second system.",
    "The remaining genera were less transferable. Mortierella increased in all three first-system contrasts but decreased in both second-system contrasts. Terrimonas increased in the first system but was near zero or decreased in the second. Chitinophaga varied with stage and system, and Botryotrichum increased overall in the first system but was rare, decreased or absent in the second. These four genera were classified as context dependent. Gemmata showed external evidence opposite to the locally pre-specified decrease. Alternaria was sparsely detected and was classified as having insufficient external evidence. Because the apparent triplicates in the second system were operational splits of one composite biological sample, they were not used for inferential testing."
    ]),
    ("3.7. Soil properties and the white-mold points provided exploratory context", [
    "Several candidate–soil pairs had large Spearman coefficients, including Morchella with total nitrogen (ρ = 0.850), Alternaria with total phosphorus (ρ = 0.817), Terrimonas with pH (ρ = −0.783) and Gemmata with available nitrogen (ρ = 0.750). None of the bacterial or fungal genus–soil associations remained significant after Benjamini–Hochberg adjustment (all q > 0.05; Fig. S1a,b). These data therefore did not identify independent physicochemical drivers.",
    "Both markers showed lower mean alpha diversity at the white-mold-affected point than at the healthy point. Mean bacterial observed ASVs were lower by 48 and Shannon diversity by 0.101; mean fungal observed ASVs were lower by 105 and Shannon diversity by 0.583 (Fig. S1c,d). Because each condition represented one spatial point within one greenhouse, these differences describe a local case and cannot be generalized as disease-associated microbiome signatures."
    ])]
    for h, paras in sections:
        heading(doc, h, 2)
        for x in paras: p(doc, x, italic_terms=("Morchella", "Mortierella", "Alternaria", "Botryotrichum", "Terrimonas", "Chitinophaga", "Gemmata", "Pseudarthrobacter"))

def add_candidate_table(doc):
    heading(doc, "Table 1. Frozen candidate genera and external evidence classification", 2)
    data = [
    ["Domain", "Genus", "M mean (%)", "A mean (%)", "Pairs matching direction", "External grade"],
    ["Fungi", "Morchella", "29.900", "Not detected", "9/9 decrease", "Partially supported"],
    ["Fungi", "Mortierella", "2.410", "15.460", "9/9 increase", "Context dependent"],
    ["Fungi", "Alternaria", "0.353", "0.020", "9/9 decrease", "Insufficient evidence"],
    ["Fungi", "Botryotrichum", "2.250", "15.260", "7/9 increase", "Context dependent"],
    ["Bacteria", "Terrimonas", "0.467", "1.574", "9/9 increase", "Context dependent"],
    ["Bacteria", "Chitinophaga", "0.051", "0.989", "9/9 increase", "Context dependent"],
    ["Bacteria", "Gemmata", "0.564", "0.132", "8/9 decrease", "Externally opposed"],
    ["Bacteria", "Pseudarthrobacter", "8.955", "0.159", "9/9 decrease", "Partially supported"]]
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for i,row in enumerate(data):
        for j,val in enumerate(row):
            cell = table.cell(i,j); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]; para.paragraph_format.line_spacing = 1
            rr = para.add_run(val); set_font(rr, size=8.5, bold=(i==0), italic=(j==1 and i>0))
    p(doc, "Notes: M, before cultivation; A, after cultivation. Values are group mean relative abundances. Candidate grades describe external directional evidence and are not biomarker-performance categories.")

def add_discussion(doc):
    heading(doc, "4. Discussion", 1)
    sections = [
    ("4.1. Community-level restructuring was the most transferable result", [
    "The strongest finding was not a universally directional genus but a coordinated reorganization of bacterial and fungal community structure. The paired design reduced variation arising from different baseline locations, bacterial results were stable across distance metrics and rarefaction depths, and fungal effect sizes were similar for all ITS ASVs and the UNITE-assigned fungal subset. Broad phylum-level replacement further indicated that the beta-diversity signal was not produced only by a few rare ASVs. These converging results support cultivation-stage-associated community restructuring. They do not, however, isolate cultivation as the sole cause because M and A samples were sequenced in different batches and the nine pairs were nested within only three greenhouses.",
    "The correspondence between bacterial and fungal distance matrices suggests that the two microbial domains shared major sample-level gradients. Common responses to pH, nutrient availability, morel mycelial development, exogenous nutrient inputs and residual organic matter are plausible explanations. Previous morel studies have likewise reported parallel bacterial and fungal succession through cultivation stages and under continuous cropping (Longley et al., 2019; Tan et al., 2021; Liu et al., 2022; Zhang et al., 2023; Yue et al., 2024). Yet part of the observed Mantel relationship may arise from the shared M/A grouping and greenhouse structure. The absence of correlation between paired bacterial and fungal alpha-diversity changes is therefore informative: two domains can occupy similar positions along a compositional gradient without gaining or losing richness and evenness in parallel. Cross-domain coordination should be claimed at the beta-structure level, not as a unified diversity mechanism."
    ]),
    ("4.2. Compositional replacement provides an ecological setting, not absolute population change", [
    "After cultivation, Patescibacteria and Bacteroidota occupied larger fractions of the bacterial community, whereas Actinomycetota declined sharply. In the fungal dataset, Ascomycota declined while Mortierellomycota increased. Such patterns may reflect changing carbon substrates, fungal necromass, nutrient availability and competition during the cultivation cycle. However, amplicon relative abundance is constrained by the total compositional denominator: one lineage can increase proportionally because it proliferates, because another lineage declines or because both processes occur (Morton et al., 2019; Lin and Peddada, 2020). Without marker-specific absolute quantification, these phylum shifts cannot be translated into population-size changes.",
    "This distinction also motivated the staged candidate strategy. Abundance-weighted ranking identified visually prominent compositional changes, but candidate freezing additionally required directional consistency across paired positions and greenhouse means plus sequence-level taxonomic support. This approach does not solve the small-sample problem or create formal biomarkers. It does reduce the risk that a candidate is driven by a single abundant sample, a weak genus assignment or post hoc selection after public results are known."
    ]),
    ("4.3. Ecological interpretation of the fungal candidates", [
    "Morchella provided the clearest local decrease: its relative signal fell from 29.90% to no detection in every matched pair. Cultivation-time studies have reported early dominance of Morchella followed by replacement by other fungi (Longley et al., 2019), supporting stage-linked attenuation as an ecological context. The present ITS data cannot distinguish living mycelium from relic DNA and are sensitive to rDNA copy number and amplification efficiency. Consequently, non-detection does not mean that absolute Morchella biomass became zero. Resource depletion and niche release after mycelial development offer a testable explanation, but confirmation requires Morchella-specific quantitative PCR, RNA-based assays or direct mycelial biomass measurements.",
    "Mortierella increased in all local pairs and was supported by multiple ASVs. Isolate-based work has shown phosphate-solubilizing activity for some Mortierella strains and changes in available phosphorus and enzyme activity in defined plant–mycorrhizal systems (Zhang et al., 2011). This establishes functional potential within the genus, not the function of the ASVs detected here. None of the Mortierella–soil-property associations passed multiplicity correction. Its enrichment may therefore be discussed as consistent with a shift toward saprotrophic resource use or altered phosphorus niches, but not as evidence that Mortierella drove phosphorus mobilization in the greenhouses.",
    "Alternaria declined consistently locally but was sparsely detected in the external datasets. The genus encompasses both saprotrophs and plant pathogens, and population-level variation in greenhouse Alternaria alternata isolates has been demonstrated (Yang et al., 2022). Because the present sequences support only genus-level assignment, the decrease cannot be interpreted as removal of a morel pathogen or reduction in disease risk. Substrate replacement or competitive filtering are possible explanations requiring species-resolved and culture-based tests.",
    "Botryotrichum showed a large mean increase but was inconsistent in one greenhouse, justifying its supporting-candidate status. A cultured Botryotrichum isolate has exhibited cellulose and lignin degradation under solid-substrate conditions (Kapoor et al., 1978). That result provides a plausible link to residual plant or fungal polymers, but it cannot assign lignocellulolytic activity to the ASVs detected here. The greenhouse-specific reversal is equally important and suggests that resource availability or initial community state conditioned the response."
    ]),
    ("4.4. Ecological interpretation of the bacterial candidates", [
    "Terrimonas increased in all nine local pairs. In soil microcosms, Terrimonas became enriched following low-dose chitin addition and was described as a potential chitin-responsive lineage (Jacquiod et al., 2013). Enrichment does not prove direct chitin hydrolysis. The present pattern is therefore compatible with, but does not demonstrate, use of fungal-derived polymers or metabolites released during necromass turnover. The relatively large negative correlation with pH did not survive false-discovery-rate adjustment and cannot identify pH as the independent driver.",
    "Chitinophaga also increased in every pair. Secretome and culture experiments have shown that Chitinophaga pinensis can use plant and fungal biomass and deploy extracellular enzymes against several beta-glucans (McKee et al., 2019). This is direct evidence for a named species under controlled conditions, not functional evidence for all Chitinophaga members. A post-cultivation increase in fungal residues or complex carbohydrates is a biologically coherent hypothesis for the local enrichment, but it requires isolation, enzyme assays, metagenomics or stable-isotope substrate tracing.",
    "Gemmata decreased in eight pairs but showed an opposing direction in public data. Cultivated Gemmata-like soil organisms are aerobic chemoheterotrophs, while their detailed roles in production soils remain poorly resolved (Wang et al., 2002). Broad functional claims based only on Planctomycetota membership would therefore be unwarranted. The disagreement between local and external evidence instead identifies Gemmata as a system-sensitive lineage whose response may depend on starting soil, water regime, pH or cultivation stage.",
    "Pseudarthrobacter showed the largest bacterial decline and received same-direction evidence in the first external system. Individual Pseudarthrobacter isolates can exhibit nitrogen fixation, mineral solubilization or indole-3-acetic-acid production, and one Pseudarthrobacter chlorophenolicus isolate promoted tomato growth in a defined plant–substrate experiment (Issifu et al., 2022). Those strain- and host-specific observations do not establish that the local Pseudarthrobacter ASVs were beneficial to morels. The decline could reflect loss of particular accessible substrates or compositional dilution by expanding taxa. Absolute quantification and isolate reintroduction are required to discriminate these alternatives."
    ]),
    ("4.5. External evidence supports a transferable community background rather than universal genus biomarkers", [
    "The public-data analysis deliberately retained contradictory and non-detectable outcomes. Only Morchella and Pseudarthrobacter reached partial support; four genera changed direction among stages or systems, Gemmata was externally opposed and Alternaria lacked sufficient detection. This pattern does not invalidate the paired local observations. Instead, it shows that the direction of an individual genus is more sensitive than the community-level restructuring signal to cultivation regime, sampling stage, initial soil and technical protocol.",
    "The external evidence also had unequal design strength. Two BioProjects arose from one experimental system and cannot be counted as two independent replications. The other system contained operational splits of pooled biological material, so apparent file-level replication did not provide biological replication. We therefore use the terms external stress test and directional evidence rather than validation cohort. A definitive transferability test would prospectively sample multiple independent greenhouses per management condition, process all samples in balanced laboratory batches and quantify the frozen candidates without redefining the list after outcomes are observed."
    ]),
    ("4.6. Soil and white-mold observations define follow-up questions but not drivers", [
    "Several genus–soil correlations were large in magnitude, but none remained significant after correction. With nine pairs nested in three greenhouses, a single greenhouse trend can produce a high rank correlation. The analyses therefore generate hypotheses about nitrogen, phosphorus and pH niches rather than evidence that those variables independently drove candidate changes.",
    "The white-mold-affected point had lower bacterial and fungal alpha diversity than the healthy point. Local diversity loss is compatible with ecological disturbance, but the comparison contains only one point per condition. Treating the three subsamples at each point as independent cases would be pseudoreplication. Disease-associated community structure, candidate pathogens or predictive signatures will require multiple independent affected and healthy greenhouses, standardized disease scoring and preferably temporal sampling before visible symptoms."
    ]),
    ("4.7. Limitations and applied implications", [
    "Four limitations determine the inferential boundary. First, the nine pairs were nested within three greenhouses, and continuous-cropping duration was completely confounded with greenhouse identity. Second, M and A corresponded to different sequencing batches, so technical variation may contribute to their separation. Third, amplicon sequencing measures marker-gene composition rather than absolute abundance, viability or metabolic function. Fourth, the external projects differed in design and contained limited independent replication. These constraints prevent causal attribution, continuous-cropping-year estimation, diagnostic-biomarker claims and functional assignment to candidate genera.",
    "Within those boundaries, the study offers an applied workflow for prioritizing follow-up measurements in morel soils. Spatial pairing, dual-domain profiling, sequence-level taxonomic review, pre-external freezing and explicit retention of opposing evidence convert a conventional amplicon survey into an auditable candidate framework. Morchella and Pseudarthrobacter merit priority in prospective absolute-quantification studies, while context-dependent genera should be evaluated as condition-specific indicators rather than discarded. This evidence hierarchy can guide a balanced validation experiment without overstating what the present production-soil dataset can establish."
    ])]
    for h, paras in sections:
        heading(doc, h, 2)
        for x in paras: p(doc, x, italic_terms=("Morchella", "Mortierella", "Alternaria", "Alternaria alternata", "Botryotrichum", "Terrimonas", "Chitinophaga", "Chitinophaga pinensis", "Gemmata", "Pseudarthrobacter", "Pseudarthrobacter chlorophenolicus"))

def add_conclusion(doc):
    heading(doc, "5. Conclusions", 1)
    p(doc, "Paired full-length 16S and ITS profiling showed that morel cultivation stage was associated with restructuring of bacterial and fungal community composition. The two domains were coordinated in beta structure but not in paired alpha-diversity response. Eight sequence-audited candidate genera showed locally stable or near-stable compositional changes, yet public-data evaluation separated partial support from context dependence, opposition and insufficient evidence. Community-level reorganization was therefore more transferable than any universal genus-level direction. Prospective multi-greenhouse replication, balanced sequencing batches, absolute quantification and isolate- or function-level experiments are required before these candidates can be used as production indicators or assigned ecological roles.")

def add_declarations(doc):
    heading(doc, "Declarations", 1)
    for h,t in [
        ("CRediT authorship contribution statement", "[AUTHOR INPUT NEEDED: insert author names and contributions using the CRediT taxonomy, including conceptualization, methodology, investigation, formal analysis, data curation, writing, visualization, supervision, project administration and funding acquisition.]"),
        ("Funding", "[AUTHOR INPUT NEEDED: funding agency, grant title and grant number. If none, state: This research received no external funding.]"),
        ("Declaration of competing interest", "The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper. [Confirm with all authors before submission.]"),
        ("Acknowledgements", "[AUTHOR INPUT NEEDED: acknowledge the sequencing provider, greenhouse managers, sample collectors and any analytical support not meeting authorship criteria.]"),
        ("Data availability", "Raw amplicon reads: [repository and accession pending]. Analysis scripts, sample metadata, source data and frozen-candidate evidence: [permanent repository DOI/URL pending]. Public datasets reanalysed in this study are available under BioProject accessions PRJNA935967, PRJNA993383 and PRJNA822526."),
        ("Declaration of generative AI and AI-assisted technologies in the manuscript preparation process", "During preparation of this work, the authors used OpenAI Codex for language drafting, structural editing and consistency checking. The authors reviewed and edited all content, verified the analyses, results and references against the underlying files and take full responsibility for the content of the publication. [Retain, revise or remove only after checking the journal policy and documenting actual use at submission.]"),
        ("Ethics statement", "Not applicable. The study analysed agricultural soil and did not involve humans, vertebrate animals or regulated clinical material. [Confirm whether local permits or landowner permissions require disclosure.]")]:
        heading(doc, h, 2); p(doc,t)

def add_refs(doc):
    heading(doc, "References", 1)
    refs = [
    "Callahan, B.J., McMurdie, P.J., Rosen, M.J., Han, A.W., Johnson, A.J.A., Holmes, S.P., 2016. DADA2: High-resolution sample inference from Illumina amplicon data. Nat. Methods 13, 581–583. https://doi.org/10.1038/nmeth.3869.",
    "Issifu, M., Songoro, E.K., Onguso, J., Ateka, E.M., Ngumi, V.W., 2022. Potential of Pseudarthrobacter chlorophenolicus BF2P4-5 as a biofertilizer for the growth promotion of tomato plants. Bacteria 1, 191–206. https://doi.org/10.3390/bacteria1040015.",
    "Jacquiod, S., Franqueville, L., Cécillon, S., Vogel, T.M., Simonet, P., 2013. Soil bacterial community shifts after chitin enrichment: an integrative metagenomic approach. PLoS ONE 8, e79699. https://doi.org/10.1371/journal.pone.0079699.",
    "Kapoor, K.K., Jain, M.K., Mishra, M.M., Singh, C.P., 1978. Cellulase activity, degradation of cellulose and lignin and humus formation by cellulolytic fungi. Ann. Microbiol. (Paris) 129B, 613–620.",
    "Knight, R., Vrbanac, A., Taylor, B.C., et al., 2018. Best practices for analysing microbiomes. Nat. Rev. Microbiol. 16, 410–422. https://doi.org/10.1038/s41579-018-0029-9.",
    "Lin, H., Peddada, S.D., 2020. Analysis of microbial compositions: a review of normalization and differential abundance analysis. npj Biofilms Microbiomes 6, 60. https://doi.org/10.1038/s41522-020-00160-w.",
    "Liu, W.-Y., et al., 2022. Determining why continuous cropping reduces the production of the morel Morchella sextelata. Front. Microbiol. 13, 903983. https://doi.org/10.3389/fmicb.2022.903983.",
    "Longley, R., Benucci, G.M.N., Mills, G., Bonito, G., 2019. Fungal and bacterial community dynamics in substrates during the cultivation of morels (Morchella rufobrunnea) indoors. FEMS Microbiol. Lett. 366, fnz215. https://doi.org/10.1093/femsle/fnz215.",
    "McKee, L.S., Martínez-Abad, A., Ruthes, A.C., Vilaplana, F., Brumer, H., 2019. Focused metabolism of β-glucans by the soil Bacteroidetes species Chitinophaga pinensis. Appl. Environ. Microbiol. 85, e02231-18. https://doi.org/10.1128/AEM.02231-18.",
    "Morton, J.T., Marotz, C., Washburne, A., et al., 2019. Establishing microbial composition measurement standards with reference frames. Nat. Commun. 10, 2719. https://doi.org/10.1038/s41467-019-10656-5.",
    "Nearing, J.T., Douglas, G.M., Hayes, M.G., et al., 2022. Microbiome differential abundance methods produce different results across 38 datasets. Nat. Commun. 13, 342. https://doi.org/10.1038/s41467-022-28034-z.",
    "Tan, H., et al., 2021. Build your own mushroom soil: microbiota succession and nutritional accumulation in semi-synthetic substratum drive the fructification of a soil-saprotrophic morel. Front. Microbiol. 12, 656656. https://doi.org/10.3389/fmicb.2021.656656.",
    "Wang, J., Jenkins, C., Webb, R.I., Fuerst, J.A., 2002. Isolation of Gemmata-like and Isosphaera-like planctomycete bacteria from soil and freshwater. Appl. Environ. Microbiol. 68, 417–422. https://doi.org/10.1128/AEM.68.1.417-422.2002.",
    "Yang, G., et al., 2022. Genetic structure and triazole antifungal susceptibilities of Alternaria alternata from greenhouses in Kunming, China. Microbiol. Spectr. 10, e00382-22. https://doi.org/10.1128/spectrum.00382-22.",
    "Yue, Y., et al., 2024. Dynamics of the soil microbial community associated with Morchella cultivation: diversity, assembly mechanism and yield prediction. Front. Microbiol. 15, 1345231. https://doi.org/10.3389/fmicb.2024.1345231.",
    "Zhang, C., Shi, X., Zhang, J., Zhang, Y., Wang, W., 2023. Dynamics of soil microbiome throughout the cultivation life cycle of morel (Morchella sextelata). Front. Microbiol. 14, 979835. https://doi.org/10.3389/fmicb.2023.979835.",
    "Zhang, H., Wu, X., Li, G., et al., 2011. Interactions between arbuscular mycorrhizal fungi and phosphate-solubilizing fungus (Mortierella sp.) and their effects on Kostelelzkya virginica growth and enzyme activities of rhizosphere and bulk soils at different salinities. Biol. Fertil. Soils 47, 543–554. https://doi.org/10.1007/s00374-011-0563-3."
    ]
    for ref in refs:
        para=p(doc, ref); para.paragraph_format.first_line_indent=Inches(-0.25); para.paragraph_format.left_indent=Inches(0.25)

def add_legends(doc):
    doc.add_page_break(); heading(doc, "Figures and figure legends", 1)
    legends = [
    ("Fig. 1. Study design, sample hierarchy and sequence-data retention.", "(a) Analysis workflow linking paired before-cultivation (M) and after-cultivation (A) soils to marker-specific analyses, cross-domain comparison, candidate freezing and public-data evaluation. (b) Nine matched spatial positions nested within three greenhouses, with three positions per greenhouse. The healthy/white-mold comparison comprised two points in one greenhouse, each represented by three spatial subsamples. (c) Retention of 16S and ITS reads through primer removal, filtering, denoising and chimera removal. (d) Samples, ASVs and reads entering downstream analyses. No sample was removed on the basis of study outcome."),
    ("Fig. 2. Paired restructuring of bacterial and fungal communities between cultivation stages.", "(a,b) Principal coordinates analyses based on Bray–Curtis dissimilarity for 16S (a) and ITS (b). Lines connect M and A soils from the same spatial position; colours indicate greenhouses. R² denotes the variance associated with cultivation stage in restricted PERMANOVA and is reported as an effect size. (c,d) Paired observed ASVs and Shannon diversity for 16S (c) and fungal ITS (d). Each marker contains nine biological pairs nested within three greenhouses."),
    ("Fig. 3. Compositional replacement of dominant bacterial and fungal taxa.", "(a,b) Sample-resolved heat maps of phylum-level relative abundance for bacteria (a) and fungi (b). Columns are individual biological soil samples ordered as matched M–A pairs within each spatial position; rows are dominant phyla, with remaining taxa grouped as Other. Colour intensity denotes relative abundance. (c,d) Eight highest abundance-weighted genus shifts for bacteria (c) and fungi (d). The x-axis shows the log2 ratio of A to M group means; point area represents the larger stage mean relative abundance. Pink and blue denote decreases and increases after cultivation, respectively. Genera were detected in at least nine M/A samples and ranked by |log2 ratio| × larger mean relative abundance. Panels c and d are descriptive and do not denote FDR-controlled differential abundance."),
    ("Fig. 4. Cross-domain beta-structure coordination without synchronized alpha-diversity responses.", "(a) Relationship between bacterial and fungal Bray–Curtis dissimilarities for the 18 M/A samples. The displayed Mantel statistic (r = 0.590) derives from the locked rarefied matrices. (b–d) Paired A-minus-M changes in bacterial and fungal alpha-diversity metrics across nine spatial positions. Colours indicate greenhouses. Technical splits were not treated as biological replicates."),
    ("Fig. 5. Eight candidate genera frozen after ASV-level reliability review.", "(a) Genus-level log2 mean abundance ratios for A relative to M, calculated with a fixed pseudocount of 1 × 10−5. (b) Mean relative abundance in M and A. Values below 0.001% are placed at the graphical lower limit only; source values are unchanged. (c) Proportion of nine matched pairs agreeing with the overall direction. Candidates were frozen before public datasets were evaluated."),
    ("Fig. 6. Public datasets provide graded rather than uniformly concordant evidence.", "(a) Direction classes for eight frozen genera in five public-data contrasts. Three contrasts in external system 1 came from one experimental system; two contrasts in external system 2 came from a second system. Each project was processed independently and ASV/count tables were not merged. (b) Log2 abundance ratios in each contrast; values beyond ±10 are clipped only for display. (c) Final evidence grades integrating direction, detectability, number of evidence systems and design constraints. The apparent replicates in external system 2 were operational splits of pooled samples and were not treated as independent biological replication."),
    ("Fig. S1. Exploratory soil-property and white-mold-point context.", "(a) Standardized A-minus-M changes in eight soil properties across nine matched positions; colours indicate greenhouses. (b) Spearman correlations between paired changes in the centered-log-ratio abundance of eight frozen genera and changes in soil properties (n = 9 positions nested within three greenhouses). All Benjamini–Hochberg-adjusted q values exceeded 0.05. (c,d) Observed ASVs and Shannon diversity for bacterial 16S (c) and fungal ITS (d) at one healthy and one white-mold-affected point in the same greenhouse. Each point is a spatial subsample and horizontal bars show point means. No disease-status inferential test was performed.")]
    figure_files = [
        ROOT / "figures_v3" / "Figure1_design_and_data_quality_v3.png",
        ROOT / "figures_v3" / "Figure2_paired_community_restructuring_v3.png",
        ROOT / "figures_v3" / "Figure3_community_composition_and_genus_shifts_v3.png",
        ROOT / "figures_v3" / "Figure4_cross_domain_coordination_v3.png",
        ROOT / "figures_v3" / "Figure5_frozen_candidate_genera_v3.png",
        ROOT / "figures_v3" / "Figure6_external_validation_evidence_v3.png",
        ROOT / "figures_v3" / "SupplementaryFigureS1_soil_and_white_mold_context_v3.png",
    ]
    for idx, ((title,body), fig_path) in enumerate(zip(legends, figure_files)):
        if idx:
            doc.add_page_break()
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        inline = pic.add_run().add_picture(str(fig_path), width=Inches(6.65))
        docPr = inline._inline.docPr
        docPr.set("title", title.split(".", 1)[0])
        docPr.set("descr", body)
        para=doc.add_paragraph(style="Figure legend")
        r=para.add_run(title+" "); set_font(r,size=10,bold=True)
        r=para.add_run(body); set_font(r,size=10)

def build_main():
    doc=Document(); configure(doc); add_title_page(doc); add_abstract(doc); add_intro(doc); add_methods(doc); add_results(doc); add_candidate_table(doc); add_discussion(doc); add_conclusion(doc); add_declarations(doc); add_refs(doc); add_legends(doc)
    props=doc.core_properties; props.title="Applied Soil Ecology manuscript draft"; props.author=""; props.subject="Morchella production soil microbiome"; props.keywords="soil microbiome; Morchella; 16S; ITS"
    doc.save(OUT)

def build_highlights():
    doc=Document(); configure(doc)
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=para.add_run("Highlights"); set_font(r,"Arial",16,bold=True)
    items=[
        "Paired full-length amplicons resolved bacterial–fungal soil shifts",
        "Cultivation stage explained 14–18% of community variation",
        "Bacterial and fungal beta structures changed coordinately",
        "Alpha-diversity responses were not synchronized across domains",
        "Public datasets exposed context dependence among frozen genera"]
    for item in items:
        para=doc.add_paragraph(style="List Bullet"); para.paragraph_format.line_spacing=2
        r=para.add_run(item); set_font(r)
        assert len(item) <= 85, (item,len(item))
    doc.core_properties.author=""
    doc.save(HIGHLIGHTS)

def add_cn_title(doc, text, size=16):
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(16)
    r=para.add_run(text); set_font(r,"Microsoft YaHei",size,bold=True)

def cn_heading(doc, text, level=1):
    para=doc.add_paragraph(style=f"Heading {level}")
    r=para.add_run(text); set_font(r,"Microsoft YaHei",14 if level==1 else 12,bold=True)
    return para

def cn_p(doc, text, bold_lead=None):
    para=doc.add_paragraph(); para.paragraph_format.line_spacing=1.5; para.paragraph_format.space_after=Pt(6)
    if bold_lead and text.startswith(bold_lead):
        r=para.add_run(bold_lead); set_font(r,"SimSun",11,bold=True); text=text[len(bold_lead):]
    r=para.add_run(text); set_font(r,"SimSun",11)
    return para

def cn_bullet(doc, text):
    para=doc.add_paragraph(style="List Bullet"); para.paragraph_format.line_spacing=1.4; para.paragraph_format.space_after=Pt(4)
    r=para.add_run(text); set_font(r,"SimSun",11)

def build_summary():
    doc=Document(); configure(doc)
    sec=doc.sections[0]; sec.left_margin=Inches(1); sec.right_margin=Inches(1)
    add_cn_title(doc,"羊肚菌土壤微生物组研究：导师汇报双总结")
    cn_p(doc,"项目范围：24份生产土壤样本，配对PacBio全长16S rRNA基因和ITS扩增子数据；核心设计为3个大棚内9个同点种植前后配对，另含同一大棚两个点位的白霉病探索性样本。")
    cn_heading(doc,"总结一：已经完成哪些分析，得到了什么数据",1)
    cn_heading(doc,"1. 样本设计和推断边界核查",2)
    cn_bullet(doc,"确认核心分析单位为9个空间位置的种植前（M）—种植后（A）同点配对样本，分布在3个大棚，每棚3个位置。")
    cn_bullet(doc,"确认三个大棚分别具有2、3、4年连作历史，但连作年限与大棚完全重合，因此不能把大棚差异解释为独立的连作年限效应。")
    cn_bullet(doc,"确认白霉病数据来自同一大棚的一个健康点和一个发病点，每点3份为空间子样，不是3个独立健康/发病重复，因此只作点位案例描述。")
    cn_bullet(doc,"识别出M/A与测序批次重合这一限制，全文统一使用‘种植阶段相关’而不使用确定性因果措辞。")
    cn_heading(doc,"2. 原始数据完整性和引物处理",2)
    cn_bullet(doc,"对48份FASTQ文件（24个样本×16S和ITS）完成MD5一致性检查，全部通过。")
    cn_bullet(doc,"核对并锁定本项目实际16S和ITS引物序列，使用Cutadapt 5.2进行双端引物识别、方向校正与切除。")
    cn_bullet(doc,"去引物后16S保留99.75%的输入序列，ITS保留99.44%，说明原始序列与引物体系匹配良好。")
    cn_heading(doc,"3. PacBio全长扩增子ASV构建与分类注释",2)
    cn_bullet(doc,"使用R 4.4.2和DADA2 1.34.0，采用PacBio误差模型、pseudo-pooling、分批误差学习和consensus去嵌合，分别构建16S与ITS ASV表。")
    cn_bullet(doc,"ITS最终获得8,056个非嵌合ASV和942,492条序列，24个样本全部保留；其中2,510个ASV被UNITE v10注释为真菌，覆盖76.77%的ITS序列。")
    cn_bullet(doc,"16S分类前获得9,752个ASV和193,540条序列；去除叶绿体、线粒体后保留9,477个ASV和188,611条序列，序列保留率97.45%。")
    cn_bullet(doc,"16S使用SILVA 138.2注释，ITS使用UNITE v10 BLAST注释，并针对候选属复核bootstrap、序列identity、coverage和优势ASV构成。")
    cn_heading(doc,"4. 多样性、群落结构和敏感性分析",2)
    cn_bullet(doc,"完成16S和ITS的Observed ASV、Shannon指数、Bray–Curtis距离、Jaccard距离、PCoA、配对受限PERMANOVA和PERMDISP分析。")
    cn_bullet(doc,"16S在1,500 reads深度下，种植阶段解释15.2%的Bray–Curtis差异（R²=0.152，探索性P=0.0039）；Jaccard结果方向一致（R²=0.117）。")
    cn_bullet(doc,"16S在1,800和2,500 reads深度下效应量仍为R²=0.162和0.176，证明结论不依赖单一稀释深度；PERMDISP未显示组间离散度差异。")
    cn_bullet(doc,"ITS全部ASV和真菌子集的M/A效应量相近，R²约为0.14–0.15；两套ITS距离矩阵的相关达到ρ=0.9852。")
    cn_bullet(doc,"Alpha多样性变化相对温和，细菌Observed ASV平均增加53.8、Shannon平均增加0.206；真菌Shannon平均增加0.942，但不同大棚间存在变异。")
    cn_heading(doc,"5. 优势类群组成变化",2)
    cn_bullet(doc,"细菌门水平：Patescibacteria由12.22%升至28.64%，Bacteroidota由11.70%升至18.41%；Actinomycetota由18.28%降至1.85%，Planctomycetota由13.32%降至9.63%。")
    cn_bullet(doc,"真菌门水平：Ascomycota由65.90%降至49.61%，Mortierellomycota由2.87%升至10.91%。")
    cn_bullet(doc,"完成门水平堆叠图和属水平丰度加权变化图，用于展示主要组成替换；该部分明确作为描述性全景，不冒充FDR显著差异属名单。")
    cn_heading(doc,"6. 细菌—真菌跨域联合分析",2)
    cn_bullet(doc,"全24样本的16S与ITS Bray–Curtis距离矩阵显著相关（Mantel r=0.587，P=0.0001）；仅保留18个M/A样本时仍为r=0.590。")
    cn_bullet(doc,"但是两种标记的配对Alpha变化不同步：Shannon变化ρ=−0.167，Observed ASV变化ρ=−0.159，均不显著。")
    cn_bullet(doc,"因此得到文章最关键的跨域结果：细菌和真菌在群落整体结构上协同重组，但各自内部丰富度/均匀度并不同步变化。")
    cn_heading(doc,"7. 8个候选属筛选、分类复核和冻结",2)
    cn_bullet(doc,"在查看公共数据结果前，综合CLR效应、9个配对方向、3个大棚平均方向、相对丰度、ASV构成和分类可靠性，冻结7个一级候选和1个支持候选。")
    cn_bullet(doc,"真菌候选：Morchella 29.90%降至未检出（9/9下降）；Mortierella 2.41%升至15.46%（9/9上升）；Alternaria 0.353%降至0.020%（9/9下降）；Botryotrichum 2.25%升至15.26%，但仅7/9上升且一个大棚方向相反。")
    cn_bullet(doc,"细菌候选：Terrimonas 0.467%升至1.574%（9/9上升）；Chitinophaga 0.051%升至0.989%（9/9上升）；Gemmata 0.564%降至0.132%（8/9下降）；Pseudarthrobacter 8.955%降至0.159%（9/9下降）。")
    cn_bullet(doc,"16S候选属reads加权bootstrap为96.51%–100%；ITS候选加权identity为97.51%–99.84%，支持属水平解释，但没有把证据不足的序列强行解释到物种。")
    cn_heading(doc,"8. 公共数据库外部方向验证",2)
    cn_bullet(doc,"独立处理PRJNA935967、PRJNA993383和PRJNA822526三个BioProject；各项目独立质控、DADA2和分类注释，没有合并不同项目的ASV或原始计数。")
    cn_bullet(doc,"PRJNA935967和PRJNA993383属于同一个外部实验体系，只计为一个证据系统；PRJNA822526为第二个体系，但组内文件是混合样本的操作拆分，生物学n=1，不计算推断性P值。")
    cn_bullet(doc,"最终证据分级：Morchella和Pseudarthrobacter为部分支持；Mortierella、Botryotrichum、Terrimonas和Chitinophaga为情境依赖；Gemmata受到外部反向证据；Alternaria外部证据不足。")
    cn_heading(doc,"9. 土壤指标和白霉病探索性分析",2)
    cn_bullet(doc,"完成8个候选属CLR变化与8项土壤指标变化的Spearman相关和BH校正。虽然Morchella–全氮、Alternaria–全磷、Terrimonas–pH等相关系数较高，但全部q>0.05。")
    cn_bullet(doc,"白霉病点位的细菌和真菌Alpha多样性均低于健康点，但只代表同一大棚的两个局部点位，不能推广为普遍的白霉病菌群特征。")
    cn_heading(doc,"10. 已形成的数据和论文材料",2)
    cn_bullet(doc,"已形成清洗后的16S/ITS ASV表、分类注释表、样本信息表、多样性矩阵、PCoA坐标、候选属证据表、土壤相关表、公共数据库外部证据矩阵以及完整运行日志和sessionInfo。")
    cn_bullet(doc,"已完成6张正文主图、1张补充图、Table 1、英文完整论文初稿、Highlights及图注，并保留可复现脚本和冻结证据目录。")

    cn_heading(doc,"总结二：论文讲了什么，得到什么结论",1)
    cn_heading(doc,"1. 论文要解决的核心问题",2)
    cn_p(doc,"这篇论文不是要证明某个菌一定导致羊肚菌增产、减产或发病，也不是建立连作年限模型。论文真正回答三个问题：第一，在同一采样位置种植前后，细菌和真菌群落是否同时发生重组；第二，两个微生物域的变化是否在群落结构和Alpha多样性两个层面都同步；第三，自有数据中方向稳定的候选属，能否在公共数据的不同栽培体系中保持相同方向。")
    cn_heading(doc,"2. 论文的主要发现",2)
    cn_bullet(doc,"羊肚菌种植前后，细菌和真菌群落均出现较明确的整体结构重排，效应量约为14%–18%，并在不同分析设置下保持稳定。")
    cn_bullet(doc,"细菌和真菌的样本间Beta结构具有较强协调性，说明两个微生物域可能共同响应种植过程中的土壤资源、菌丝生长、残体和管理环境变化。")
    cn_bullet(doc,"这种协调并不等于两个微生物域的Alpha多样性同步变化。细菌丰富度增加时，真菌丰富度不一定增加，说明跨域协同主要发生在群落组成结构，而不是简单的‘多样性一起升高或降低’。")
    cn_bullet(doc,"种植后优势类群发生明显替换，说明变化不是由少数低丰度ASV造成，而是涉及较广泛的土壤微生物群落重组。")
    cn_bullet(doc,"8个候选属在自有数据中具有较强或近稳定的配对方向，但公共数据没有支持它们全部跨体系保持一致。")
    cn_bullet(doc,"公共验证表明，Morchella和Pseudarthrobacter具有一定跨体系共性；多数属受栽培阶段、初始土壤、管理制度和技术平台影响，具有明显情境依赖性。")
    cn_heading(doc,"3. 论文的中心结论",2)
    cn_p(doc,"最核心的结论是：羊肚菌种植土壤存在可以在不同分析中识别的细菌—真菌群落重构背景，但‘群落整体发生变化’比‘某一个属永远按同一方向变化’更具有可迁移性。换句话说，羊肚菌生产土壤的微生物响应具有群落层面的共性，却没有证据支持一个普遍适用于所有地点和栽培体系的单属生物标志物。")
    cn_heading(doc,"4. 这篇论文的创新和价值",2)
    cn_bullet(doc,"使用同点前后配对设计，减少了不同采样地点起始土壤差异对结果的干扰。")
    cn_bullet(doc,"对同一批土壤同时分析全长16S和ITS，不再把细菌和真菌完全割裂，而是比较跨域群落协调。")
    cn_bullet(doc,"候选属在公共数据分析前提前冻结，并进行ASV层面的分类可靠性复核，减少事后挑选‘验证成功’菌属的偏差。")
    cn_bullet(doc,"公共数据库部分不仅报告支持结果，也保留反向和证据不足结果，从而更真实地界定候选属的适用范围。")
    cn_bullet(doc,"对于连作年限、测序批次、白霉病伪重复和相对丰度等问题主动限定结论，增强论文的可信度和抗审稿质疑能力。")
    cn_heading(doc,"5. 不能向导师或审稿人过度表述的内容",2)
    cn_bullet(doc,"不能说已经证明羊肚菌种植‘导致’这些群落变化，只能说与种植阶段相关，因为M/A和测序批次重合。")
    cn_bullet(doc,"不能比较2、3、4年连作的独立年限效应，因为每个年限只对应一个大棚。")
    cn_bullet(doc,"不能把8个候选属直接叫作促生菌、病原菌、拮抗菌或诊断标志物；目前只证明属水平相对组成变化。")
    cn_bullet(doc,"不能把白霉病6个样本当作3个健康重复和3个发病重复；它们实际是两个点位的空间子样。")
    cn_bullet(doc,"不能由扩增子相对丰度推断绝对菌量、活性或代谢功能。")
    cn_heading(doc,"6. 给导师汇报时可直接使用的一段话",2)
    cn_p(doc,"我们已经完成24份羊肚菌生产土壤的PacBio全长16S和ITS联合分析。核心数据是3个大棚、9个采样位置的种植前后同点配对。结果显示，种植阶段与细菌和真菌群落整体结构重组明显相关，两个微生物域的Beta结构变化具有较强协调性，但Alpha多样性变化并不同步。我们在公共数据分析前复核并冻结了8个候选属，再用3个公共BioProject进行外部方向验证。公共结果只对Morchella和Pseudarthrobacter提供部分支持，多数候选表现出体系和阶段依赖性。因此论文的主要结论不是找到一个普遍的有益菌或病害标志物，而是证明羊肚菌种植土壤存在跨细菌—真菌的群落重构，同时指出单个菌属的响应不能脱离具体栽培环境进行普遍化解释。")
    cn_heading(doc,"7. 下一步最值得补充的工作",2)
    cn_bullet(doc,"先向测序公司补齐PCR体系、循环程序、建库试剂、PacBio仪器和CCS参数。")
    cn_bullet(doc,"补齐采样日期、地点、土层深度、样品保存方法和8项土壤指标的具体检测方法。")
    cn_bullet(doc,"完成原始序列和分析代码的公共数据库上传，取得正式accession和仓库链接。")
    cn_bullet(doc,"如果后续能够增加实验，优先做Morchella和Pseudarthrobacter绝对定量，并增加独立大棚重复；这比继续增加常规微生物组图更能提高论文层次。")
    doc.core_properties.title="羊肚菌土壤微生物组导师汇报双总结"; doc.core_properties.author=""
    doc.save(SUMMARY)

if __name__ == "__main__":
    build_main(); build_highlights(); build_summary(); print(OUT); print(HIGHLIGHTS); print(SUMMARY)

